from __future__ import annotations

import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook

archive = Path('/home/ubuntu/upload/user_data_export_2026-08-29-11-36-55_39948119-c460-40ea-9591-d4ec11795ebc_Uskb4.zip')
output = Path('/home/ubuntu/cognitive-workspace/docs/export_inspection.txt')

with zipfile.ZipFile(archive) as zf:
    names = zf.namelist()
    xlsx_name = next(name for name in names if name.endswith('.xlsx'))
    json_name = next(name for name in names if name.endswith('.json'))
    xlsx_path = Path('/tmp') / Path(xlsx_name).name
    json_path = Path('/tmp') / Path(json_name).name
    xlsx_path.write_bytes(zf.read(xlsx_name))
    json_path.write_bytes(zf.read(json_name))

lines = []
wb = load_workbook(xlsx_path, read_only=True, data_only=True)
lines.append(f'workbook_sheets={wb.sheetnames}')
for ws in wb.worksheets:
    rows = ws.iter_rows(values_only=True)
    header = next(rows, ())
    sample = next(rows, ())
    lines.append(f'sheet={ws.title} max_row={ws.max_row} max_column={ws.max_column}')
    lines.append(f'header={list(header)}')
    lines.append(f'first_row_value_types={[type(v).__name__ for v in sample]}')

with json_path.open('r', encoding='utf-8') as fh:
    data = json.load(fh)
lines.append(f'json_top_level_type={type(data).__name__}')
if isinstance(data, list):
    lines.append(f'json_length={len(data)}')
    if data:
        first = data[0]
        lines.append(f'json_first_item_type={type(first).__name__}')
        if isinstance(first, dict):
            lines.append(f'json_first_item_keys={sorted(first.keys())}')
            for key, value in first.items():
                if isinstance(value, list):
                    lines.append(f'json_first_item_list_field={key} length={len(value)}')
                    if value and isinstance(value[0], dict):
                        lines.append(f'json_first_item_list_field_keys={key}:{sorted(value[0].keys())}')
                elif isinstance(value, dict):
                    lines.append(f'json_first_item_dict_field={key} keys={sorted(value.keys())}')
elif isinstance(data, dict):
    lines.append(f'json_top_level_keys={sorted(data.keys())}')
    for key, value in data.items():
        if isinstance(value, list):
            lines.append(f'json_top_level_list_field={key} length={len(value)}')
            if value and isinstance(value[0], dict):
                lines.append(f'json_top_level_list_field_keys={key}:{sorted(value[0].keys())}')
                for nested_key, nested_value in value[0].items():
                    if isinstance(nested_value, list):
                        lines.append(f'json_nested_list_field={key}.{nested_key} length_first_item={len(nested_value)}')
                        if nested_value and isinstance(nested_value[0], dict):
                            lines.append(f'json_nested_list_field_keys={key}.{nested_key}:{sorted(nested_value[0].keys())}')
                    elif isinstance(nested_value, dict):
                        lines.append(f'json_nested_dict_field={key}.{nested_key} keys={sorted(nested_value.keys())}')

output.write_text('\n'.join(lines) + '\n', encoding='utf-8')
print(output)
